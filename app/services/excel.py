"""당직표 엑셀 파싱.

**파일을 디스크에 쓰지 않는다.** 넘어온 바이트를 메모리에서 열어 읽고 즉시 버린다.

**시트를 통째로 메모리에 올리지 않는다.** `read_only=True` 로 행을 흘려 읽으면서
상한(행·열)까지만 격자로 만든다. 10MB 짜리 xlsx 가 압축 해제 시 수 GB 로 부풀어도
읽는 양이 상한에서 끊긴다.

시트 이름을 고정하지 않는다. 원본이 「26년 당직 일정_정리」라 해서 그대로 박아 두면
내년 파일(「27년 …」)에서 바로 깨진다. 대신 **머리글 '날짜'를 찾아 표의 위치를 스스로 알아낸다.**

표는 두 종류다.
  · 주말/휴일 표 — 월 · 날짜 · 요일 · 구분 · 조직 · 담당자 · 비고
  · 평일 표     — 월 · 날짜 · 요일 · 구분 · 담당자
'날짜' 머리글 오른쪽 세 칸 뒤에 '조직'이 있으면 주말 표, 없으면 평일 표로 판별한다.
"""

from __future__ import annotations

import io
import logging
from datetime import date, datetime
from typing import Any, Sequence

from openpyxl import load_workbook

from app.errors import ProblemError
from app.models import NOTE_MAX, PERSON_MAX, clean_text

logger = logging.getLogger("duty.excel")

VALID_DAY_TYPES = {"평일", "휴일"}
VALID_ORGS = {"영기", "소강", "도강", "전산휴무"}

# 읽어들일 상한. 한 해가 366행이므로 넉넉하면서도 메모리를 묶어 둔다.
_MAX_ROWS = 2000
_MAX_COLS = 60
# 머리글을 찾을 범위
_HEADER_SEARCH_ROWS = 20

DutyRow = tuple[date, str, str | None, str, str]
Grid = list[tuple[Any, ...]]


def _cell_text(value: Any) -> str:
    """셀 값을 비교 가능한 문자열로 만든다. 머리글에 줄바꿈이 섞여 있어도 견딘다."""
    if value is None:
        return ""
    return clean_text(str(value)).replace(" ", "")


def _at(grid: Grid, row: int, col: int) -> Any:
    """격자에서 값을 꺼낸다. 범위를 벗어나면 None (행마다 길이가 다를 수 있다)."""
    if 0 <= row < len(grid):
        line = grid[row]
        if 0 <= col < len(line):
            return line[col]
    return None


def _as_date(value: Any) -> date | None:
    """셀 값을 날짜로 바꾼다. 날짜가 아니면 None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def _read_grid(sheet: Any) -> Grid:
    """시트를 상한까지만 격자로 읽는다. read_only 모드라 행을 흘려 받는다."""
    grid: Grid = []
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index >= _MAX_ROWS:
            break
        grid.append(tuple(row[:_MAX_COLS]))
    return grid


def _find_tables(grid: Grid) -> list[dict[str, Any]]:
    """'날짜' 머리글을 찾아 표의 시작 위치와 열 배치를 알아낸다. (행·열은 0-기준)"""
    tables: list[dict[str, Any]] = []

    for row in range(min(len(grid), _HEADER_SEARCH_ROWS)):
        for col in range(len(grid[row])):
            if _cell_text(_at(grid, row, col)) != "날짜":
                continue
            has_org = _cell_text(_at(grid, row, col + 3)) == "조직"
            tables.append(
                {
                    "header_row": row,
                    "date_col": col,
                    "type_col": col + 2,  # 구분
                    "org_col": col + 3 if has_org else None,
                    "person_col": col + 4 if has_org else col + 3,
                    "note_col": col + 5 if has_org else None,
                    "kind": "휴일표" if has_org else "평일표",
                }
            )
    return tables


def _read_table(grid: Grid, table: dict[str, Any]) -> dict[date, dict[str, Any]]:
    """표 하나를 날짜별 dict 로 읽는다. 날짜 칸이 비면 표가 끝난 것으로 본다."""
    rows: dict[date, dict[str, Any]] = {}
    row = table["header_row"] + 1

    while row < len(grid):
        raw_date = _at(grid, row, table["date_col"])
        if raw_date is None or (isinstance(raw_date, str) and not raw_date.strip()):
            break

        # 사용자에게 보이는 행 번호는 엑셀과 같게 1-기준으로 돌려준다.
        excel_row = row + 1

        duty_date = _as_date(raw_date)
        if duty_date is None:
            raise ProblemError(
                400, "엑셀 형식 오류", f"{table['kind']} {excel_row}행의 날짜를 읽을 수 없습니다."
            )

        day_type = _cell_text(_at(grid, row, table["type_col"]))
        if day_type not in VALID_DAY_TYPES:
            raise ProblemError(
                400,
                "엑셀 형식 오류",
                f"{table['kind']} {excel_row}행의 구분은 '평일' 또는 '휴일'이어야 합니다.",
            )

        org = None
        if table["org_col"] is not None:
            org_text = _cell_text(_at(grid, row, table["org_col"]))
            if org_text and org_text not in VALID_ORGS:
                raise ProblemError(
                    400,
                    "엑셀 형식 오류",
                    f"{table['kind']} {excel_row}행에 알 수 없는 조직이 있습니다: {org_text[:20]}",
                )
            org = org_text or None

        person = clean_text(str(_at(grid, row, table["person_col"]) or ""))
        note = ""
        if table["note_col"] is not None:
            note = clean_text(str(_at(grid, row, table["note_col"]) or ""))

        if len(person) > PERSON_MAX or len(note) > NOTE_MAX:
            raise ProblemError(
                400,
                "엑셀 형식 오류",
                f"{table['kind']} {excel_row}행의 담당자 또는 비고가 너무 깁니다.",
            )

        rows[duty_date] = {
            "day_type": day_type,
            "org": None if day_type == "평일" else org,
            "person": person,
            "note": note,
        }
        row += 1

    return rows


def parse_duty_workbook(content: bytes) -> list[DutyRow]:
    """엑셀 바이트를 읽어 하루 = 한 행 목록으로 만든다.

    한 행이라도 이상하면 **여기서 예외로 끝난다.** 호출자는 이 함수가 성공한 뒤에야
    DB 를 건드리므로, 형식 오류로 기존 데이터가 손상되는 일이 없다.
    """
    try:
        # data_only=True : 수식을 평가하지 않고 저장된 값만 읽는다(수식 폭탄 차단)
        # read_only=True : 시트를 통째로 메모리에 올리지 않고 행 단위로 흘려 읽는다
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        logger.warning("엑셀 열기 실패: %s", type(exc).__name__)
        raise ProblemError(400, "엑셀 형식 오류", "엑셀 파일(.xlsx)로 읽을 수 없습니다.") from None

    try:
        merged: dict[date, dict[str, Any]] = {}
        found = False

        for sheet in workbook.worksheets:
            grid = _read_grid(sheet)
            tables = _find_tables(grid)
            if not tables:
                continue
            found = True
            # 평일표를 먼저 깔고 휴일표로 덮는다 — 같은 날짜면 휴일 정보가 우선이다.
            for table in sorted(tables, key=lambda t: t["kind"] == "휴일표"):
                merged.update(_read_table(grid, table))

        if not found:
            raise ProblemError(
                400,
                "엑셀 형식 오류",
                "당직표 형식을 인식할 수 없습니다. '날짜'·'구분' 머리글이 있는 표가 필요합니다.",
            )
        if not merged:
            raise ProblemError(400, "엑셀 형식 오류", "당직 일정 행을 찾지 못했습니다.")

        return [
            (duty_date, row["day_type"], row["org"], row["person"], row["note"])
            for duty_date, row in sorted(merged.items())
        ]
    finally:
        workbook.close()
